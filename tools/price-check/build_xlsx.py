"""Excel для проверки цен TheDoubleF.

Вход — CSV из `server/scripts/tdf-price-check.ts`. Каждая проверка — формула
от цен поставщика и параметров листа «Сводка»: закупка, пошлина, затраты,
комиссия, остаток, маржа, сетка Pricing Logic v1.0 и зачёркнутая цена
(розница в Италии с НДС). Поменял параметр — пересчиталось всё.

    python build_xlsx.py tdf-price-check.csv Italian_Edit_Price_Check.xlsx [дата фида]
"""
import csv
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
feed_date = sys.argv[3] if len(sys.argv) > 3 else ""
rows = list(csv.DictReader(open(src, encoding="utf-8-sig")))
rows.sort(key=lambda r: (r["in_feed"] != "да", r["brand"].lower(), r["title"].lower()))

F = "Arial"
base, bold = Font(name=F), Font(name=F, bold=True)
blue, green = Font(name=F, color="0000FF"), Font(name=F, color="008000")
grey = Font(name=F, color="666666", italic=True)
head_fill, head_font = PatternFill("solid", start_color="1F2937"), Font(name=F, bold=True, color="FFFFFF")
red_fill = PatternFill("solid", start_color="FECACA")
amber_fill = PatternFill("solid", start_color="FDE68A")

wb = Workbook()
s = wb.active
s.title = "Сводка"
for col, w in zip("ABC", (50, 16, 78)):
    s.column_dimensions[col].width = w

s["A1"] = "Italian Edit — проверка цен TheDoubleF (Pricing Logic v1.0, рынок США)"
s["A1"].font = Font(name=F, bold=True, size=14)
s["A2"] = (f"Цены и данные — как сейчас в магазине; цены поставщика — из фида {feed_date}. "
           "Синие ячейки — исходные данные, чёрные — формулы, зелёные — итоги по листу «Цены».")
s["A2"].font = grey
s.merge_cells("A2:C2")


def header(sheet, row, labels):
    for i, label in enumerate(labels, 1):
        c = sheet.cell(row=row, column=i, value=label)
        c.font, c.fill = head_font, head_fill


header(s, 4, ["Параметр", "Значение", "Источник"])
params = [  # строки 5–14
    ("Курс EUR → USD", 1.1206, "0.0000", "ЕЦБ на 09.10.2026; на Render — TDF_EUR_USD"),
    ("Комиссия оплаты, %", 0.029, "0.0%", "Pricing Logic v1.0, «Тесты приёмки»"),
    ("Комиссия оплаты, фикс. USD", 0.30, "$#,##0.00", "Pricing Logic v1.0, «Тесты приёмки»"),
    ("Доставка, USD на отправление", 35, "$#,##0.00", "Таблица экономики (оценка до договорного тарифа)"),
    ("Таможенное оформление, USD", 25, "$#,##0.00", "Таблица экономики (оценка)"),
    ("Пошлина без HS-кода: аксессуары и сумки из ЕС", 0.15, "0.0%", "Оценка по группе — только если кода нет в тарифе США"),
    ("Пошлина без HS-кода: одежда из ЕС", 0.22, "0.0%", "Оценка по группе"),
    ("Пошлина без HS-кода: обувь из ЕС", 0.25, "0.0%", "Оценка по группе"),
    ("Пошлина: не ЕС — не ниже", 0.30, "0.0%", "Осторожная оценка до подтверждения брокером"),
    ("НДС Италии для зачёркнутой цены", 0.22, "0.0%", "Решение 13.10.2026, вариант B"),
]
for i, (k, v, fmt, note) in enumerate(params, 5):
    s.cell(row=i, column=1, value=k).font = base
    c = s.cell(row=i, column=2, value=v)
    c.font, c.number_format = blue, fmt
    s.cell(row=i, column=3, value=note).font = grey
P = {name: f"'Сводка'!$B${row}" for name, row in zip(
    ("fx", "fee", "fix", "ship", "broker", "d_acc", "d_cloth", "d_shoes", "d_noneu", "vat"), range(5, 15))}

s["A15"] = ("Пошлина по HS-коду: базовая ставка США (MFN, USITC) по 6 знакам кода поставщика; для ЕС — не ниже 15% "
            "(сделка США–ЕС с 01.07.2026), для остальных — не ниже 30%. Ставка строки — в колонке «Ставка пошлины».")
s["A15"].font = grey
s.merge_cells("A15:C15")
s["A16"] = "Сетка минимумов («Логика v1.0»): оба условия по итоговой цене"
s["A16"].font = bold
header(s, 17, ["Цена от, USD (включительно)", "Мин. маржа", "Мин. остаток, USD"])
for i, (lo, m, rem) in enumerate([(0, 0.20, 40), (1000, 0.18, 200), (2500, 0.15, 450)], 18):
    for col, v, fmt in ((1, lo, "$#,##0"), (2, m, "0%"), (3, rem, "$#,##0")):
        c = s.cell(row=i, column=col, value=v)
        c.font, c.number_format, c.alignment = blue, fmt, Alignment(horizontal="left")
G = "'Сводка'"

# ── Лист «Цены» ──────────────────────────────────────────────────────────
t = wb.create_sheet("Цены")
cols = [  # (заголовок, ширина)
    ("Бренд", 16), ("Модель", 42), ("Статус", 9), ("Раздел", 11), ("Подкатегория", 13),
    ("Артикул TDF", 16), ("HS-код", 11), ("Страна", 7), ("ЕС", 5),
    ("price, EUR", 10), ("sale_price, EUR", 10), ("Закупка, EUR", 10), ("Закупка, USD", 11),
    ("Cost per item на сайте", 11), ("Закупка сходится", 9), ("Ставка пошлины", 9), ("Пошлина, USD", 10),
    ("Прямые затраты, USD", 11), ("Наша цена на сайте", 10), ("Комиссия, USD", 9), ("Остаток, USD", 10),
    ("Маржа", 8), ("Мин. маржа", 8), ("Мин. остаток", 9), ("Сетка", 11),
    ("Зачёркнутая — ожидаемая", 11), ("Зачёркнутая на сайте", 11), ("Зачёркнутая сходится", 10),
    ("Скидка на сайте", 8), ("Режим цены", 9), ("Пометки", 24), ("Остаток, шт.", 8),
    ("Ссылка поставщика", 30), ("Админка", 30),
    ("Основание пошлины", 12), ("Farfetch US, итог USD", 11), ("Farfetch: проверено", 10), ("Farfetch: пригодно", 14),
    ("Italist US, итог USD", 11), ("Italist: проверено", 10), ("Italist: пригодно", 14),
    ("Ориентир: дешевле из пригодных", 11), ("Наша цена к ориентиру", 10), ("Маржа по цене ориентира", 10),
    ("Метод цены", 9), ("Решение v1.0", 14),
]
for i, (h, w) in enumerate(cols, 1):
    c = t.cell(row=1, column=i, value=h)
    c.font, c.fill = head_font, head_fill
    c.alignment = Alignment(wrap_text=True, vertical="center")
    t.column_dimensions[get_column_letter(i)].width = w
t.row_dimensions[1].height = 44
t.freeze_panes = "C2"


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


for i, r in enumerate(rows, 2):
    inputs = {
        1: r["brand"], 2: r["title"], 3: r["status"], 4: r["section"], 5: r["subcategory"],
        6: r["model_code"], 7: r["hs_code"], 8: r["country"], 9: r["eu"],
        10: num(r["price_eur"]), 11: num(r["sale_price_eur"]), 14: num(r["store_cost"]),
        19: num(r["store_price"]), 27: num(r["store_compare_at"]),
        30: r["price_mode"], 31: r["review"], 32: int(r["stock"] or 0),
        33: r["supplier_link"], 34: r["admin_url"],
        16: num(r["duty_rate"]), 35: r["duty_basis"],
        36: num(r["farfetch_usd"]), 37: r["farfetch_checked"], 38: r["farfetch_usable"],
        39: num(r["italist_usd"]), 40: r["italist_checked"], 41: r["italist_usable"],
        45: r["price_method"], 46: r["price_decision"],
    }
    for col, v in inputs.items():
        if v not in (None, ""):
            t.cell(row=i, column=col, value=v)
    f = {
        12: f'=IF(J{i}="","",IF(N(K{i})>0,K{i},J{i}))',
        13: f'=IF(L{i}="","",ROUND(L{i}*{P["fx"]},2))',
        15: f'=IF(OR(M{i}="",N{i}=""),"",IF(ABS(M{i}-N{i})<=0.01,"да","НЕТ"))',
        17: f'=IF(M{i}="","",ROUND(M{i}*P{i},2))',
        18: f'=IF(M{i}="","",M{i}+Q{i}+{P["ship"]}+{P["broker"]})',
        20: f'=IF(N(S{i})=0,"",ROUND(S{i}*{P["fee"]}+{P["fix"]},2))',
        21: f'=IF(OR(R{i}="",T{i}=""),"",ROUND(S{i}-R{i}-T{i},2))',
        22: f'=IF(U{i}="","",U{i}/S{i})',
        23: f'=IF(N(S{i})=0,"",IF(S{i}<{G}!$A$19,{G}!$B$18,IF(S{i}<{G}!$A$20,{G}!$B$19,{G}!$B$20)))',
        24: f'=IF(N(S{i})=0,"",IF(S{i}<{G}!$A$19,{G}!$C$18,IF(S{i}<{G}!$A$20,{G}!$C$19,{G}!$C$20)))',
        25: f'=IF(U{i}="","",IF(AND(U{i}>=X{i},U{i}>=ROUNDUP(S{i}*W{i},2)),"Проходит","Ниже порога"))',
        26: (f'=IF(OR(J{i}="",N(S{i})=0),"",IF(ROUND(ROUND(ROUND(J{i}*(1+{P["vat"]}),2)*{P["fx"]},2),0)'
             f'>S{i}*1.01,ROUND(ROUND(ROUND(J{i}*(1+{P["vat"]}),2)*{P["fx"]},2),0),""))'),
        28: f'=IF(J{i}="","",IF(N(Z{i})=N(AA{i}),"да","НЕТ"))',
        29: f'=IF(N(AA{i})>N(S{i}),1-S{i}/AA{i},"")',
        42: (f'=IF(AND(AL{i}="да",AO{i}="да"),MIN(AJ{i},AM{i}),IF(AL{i}="да",AJ{i},IF(AO{i}="да",AM{i},"")))'),
        43: f'=IF(OR(AP{i}="",N(S{i})=0),"",S{i}/AP{i}-1)',
        44: f'=IF(OR(AP{i}="",R{i}=""),"",(AP{i}-R{i}-ROUND(AP{i}*{P["fee"]}+{P["fix"]},2))/AP{i})',
    }
    for col, formula in f.items():
        t.cell(row=i, column=col, value=formula)
    for col in range(1, 47):
        c = t.cell(row=i, column=col)
        c.font = blue if col in inputs and col not in (30, 31, 33, 34, 35, 38, 41, 45, 46) else base
    for col in (10, 11, 12):
        t.cell(row=i, column=col).number_format = "€#,##0.00"
    for col in (13, 14, 17, 18, 20, 21, 24):
        t.cell(row=i, column=col).number_format = "$#,##0.00"
    for col in (19, 26, 27, 36, 39, 42):
        t.cell(row=i, column=col).number_format = "$#,##0"
    for col in (16, 22, 23, 29, 43, 44):
        t.cell(row=i, column=col).number_format = "0.0%"

last = len(rows) + 1
t.auto_filter.ref = f"A1:AT{last}"
for rng, cond in ((f"Y2:Y{last}", 'Y2="Ниже порога"'), (f"O2:O{last}", 'O2="НЕТ"'), (f"AB2:AB{last}", 'AB2="НЕТ"')):
    t.conditional_formatting.add(rng, FormulaRule(formula=[cond], fill=red_fill))
t.conditional_formatting.add(f"A2:B{last}", FormulaRule(formula=[f'$AD2="ручная"'], fill=amber_fill))
t.cell(row=last + 2, column=1, value=(
    "Синие — исходные данные: из фида (price, sale_price, страна) и из магазина (Cost per item, наша цена, "
    "зачёркнутая цена). Остальное — формулы. Жёлтые строки — цена поставлена вручную: сетку она может и не проходить. "
    "Внизу таблицы — товары, которых сегодня нет в фиде (снят поставщиком или распродан)."
)).font = grey

# ── Итоги на «Сводке» ──────────────────────────────────────────────────────
C = "'Цены'"
s["A22"] = "Итоги по магазину"
s["A22"].font = bold
header(s, 23, ["Показатель", "Значение", "Что значит"])
summary = [
    ("Товаров TheDoubleF в магазине", f"=COUNTA({C}!B2:B{last})", "0", ""),
    ("…из них сегодня нет в фиде", f'=COUNTBLANK({C}!J2:J{last})', "0", "Сняты поставщиком или распроданы — остаток 0"),
    ("Проходят оба минимума сетки", f'=COUNTIF({C}!Y2:Y{last},"Проходит")', "0", ""),
    ("Ниже порога", f'=COUNTIF({C}!Y2:Y{last},"Ниже порога")', "0", "Красные строки; у авто-цен должно быть 0"),
    ("…из них с автоматической ценой", f'=COUNTIFS({C}!Y2:Y{last},"Ниже порога",{C}!AD2:AD{last},"авто")', "0", "Должно быть 0"),
    ("Закупка не сходится с Cost per item", f'=COUNTIF({C}!O2:O{last},"НЕТ")', "0", "Поставщик сменил цену — синк обновит при следующем проходе"),
    ("Зачёркнутая цена не сходится", f'=COUNTIF({C}!AB2:AB{last},"НЕТ")', "0", "То же"),
    ("Товаров со скидкой на сайте", f"=COUNT({C}!AC2:AC{last})", "0", "Есть зачёркнутая цена"),
    ("Медианная скидка", f"=MEDIAN({C}!AC2:AC{last})", "0.0%", ""),
    ("Медианная маржа", f"=MEDIAN({C}!V2:V{last})", "0.0%", "До рекламы, возвратов и постоянных расходов"),
    ("Минимальная маржа", f"=MIN({C}!V2:V{last})", "0.0%", ""),
    ("Медианная цена", f"=MEDIAN({C}!S2:S{last})", "$#,##0", ""),
    ("С ручной ценой", f'=COUNTIF({C}!AD2:AD{last},"ручная")', "0", "Жёлтые строки"),
    ("Пошлина по HS-коду", f'=COUNTIF({C}!AI2:AI{last},"HS-код")', "0", "Остальные — оценка по группе (кода нет в фиде)"),
    ("Моделей с пригодной ценой конкурента", f"=COUNT({C}!AP2:AP{last})", "0", "USD, витрина US, не старше 14 дней — из data/tdf-market.csv"),
    ("…из них мы дешевле ориентира", f'=COUNTIF({C}!AQ2:AQ{last},"<=0")', "0", ""),
    ("…медианная маржа по цене ориентира", f"=IFERROR(MEDIAN({C}!AR2:AR{last}),\"\")", "0.0%", "Сколько бы зарабатывали, продавая по цене конкурента"),
]
for i, (k, formula, fmt, note) in enumerate(summary, 24):
    s.cell(row=i, column=1, value=k).font = base
    c = s.cell(row=i, column=2, value=formula)
    c.font, c.number_format, c.alignment = green, fmt, Alignment(horizontal="left")
    s.cell(row=i, column=3, value=note).font = grey

r0 = 24 + len(summary) + 1
s.cell(row=r0, column=1, value="Как проверить товар вручную").font = bold
steps = [
    "1. Найти модель на vip.thedoublef.com по «Артикулу TDF» (или по ссылке поставщика), сверить price и sale_price (EUR без НДС).",
    "2. Закупка = sale_price (если 0 — price) × курс. Должна совпасть с Cost per item в админке Shopify.",
    "3. Затраты = закупка + пошлина (ставка по разделу и стране) + доставка + оформление; комиссия = цена × 2,9% + $0,30.",
    "4. Маржа = (цена − затраты − комиссия) / цена; оба минимума сетки по итоговой цене. Наша цена — минимальная, проходящая сетку, в целых долларах.",
    "5. Зачёркнутая цена = price × 1,22 × курс до доллара; на сайте — только если выше нашей цены больше чем на 1%.",
    "6. Farfetch / Italist сравнивать только в долларах и на версии для США (VPN с американским IP): сайты переключают регион и валюту сами.",
    "7. Цены конкурентов вносить в server/data/tdf-market.csv (или Google-таблицу TDF_MARKET_CSV_URL): артикул, farfetch/italist, итог с доставкой и пошлиной в USD, USD, US, дата.",
]
for i, text in enumerate(steps, r0 + 1):
    c = s.cell(row=i, column=1, value=text)
    c.font, c.alignment = base, Alignment(wrap_text=True, vertical="top")
    s.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
    s.row_dimensions[i].height = 30

wb.calculation.fullCalcOnLoad = True
wb.save(out)
print(f"{len(rows)} строк → {out}")
